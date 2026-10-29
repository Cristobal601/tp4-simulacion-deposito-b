"""Exporta el vector de estado a una planilla de Excel (.xlsx)."""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from simulacion.modelo import CLAVES, COLUMNAS


def exportar_excel(ruta, filas):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Vector de estado"

    # Fila 1: grupos (celdas combinadas) - Fila 2: titulos de cada columna
    columna = 1
    for grupo, columnas in COLUMNAS:
        hoja.cell(row=1, column=columna, value=grupo).font = Font(bold=True)
        hoja.merge_cells(start_row=1, start_column=columna,
                         end_row=1, end_column=columna + len(columnas) - 1)
        for clave, titulo in columnas:
            hoja.cell(row=2, column=columna, value=titulo).font = Font(bold=True)
            columna += 1

    # Datos
    for numero_fila, fila in enumerate(filas, start=3):
        for numero_columna, clave in enumerate(CLAVES, start=1):
            hoja.cell(row=numero_fila, column=numero_columna, value=fila.get(clave))

    for celda in hoja[1]:
        celda.alignment = Alignment(horizontal="center")
    hoja.freeze_panes = "A3"   # encabezados fijos al hacer scroll
    libro.save(ruta)
